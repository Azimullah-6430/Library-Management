"""
Excel Manager
-------------
Handles all Excel file operations for the Library Book Management System.
"""

import os
from openpyxl import Workbook, load_workbook


# Excel column headers
HEADERS = [
    "Book ID",
    "Book Title",
    "Author",
    "ISBN",
    "Category",
    "Publisher",
    "Publication Year",
    "Total Copies",
    "Available Copies",
    "Shelf Number",
    "Cover Image",
    "Date Added"
]


def create_excel_file(excel_file):
    """
    Create the Excel file and Books worksheet if they do not exist.
    """

    # Create parent directory if it does not exist
    folder = os.path.dirname(excel_file)

    if folder:
        os.makedirs(folder, exist_ok=True)

    # Do not recreate an existing file
    if os.path.exists(excel_file):
        return

    # Create workbook
    workbook = Workbook()

    # Rename default worksheet
    sheet = workbook.active
    sheet.title = "Books"

    # Add headers
    sheet.append(HEADERS)

    # Make headers bold
    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True)

    # Set useful column widths
    column_widths = {
        "A": 15,
        "B": 30,
        "C": 25,
        "D": 20,
        "E": 20,
        "F": 25,
        "G": 20,
        "H": 15,
        "I": 18,
        "J": 18,
        "K": 35,
        "L": 20
    }

    for column, width in column_widths.items():
        sheet.column_dimensions[column].width = width

    # Freeze the header row
    sheet.freeze_panes = "A2"

    # Save the workbook
    workbook.save(excel_file)


def get_all_books(excel_file):
    """
    Read all books from the Excel file.

    Returns:
        list: List of dictionaries containing book information.
    """

    create_excel_file(excel_file)

    workbook = load_workbook(excel_file, data_only=True)
    sheet = workbook["Books"]

    books = []

    for row in sheet.iter_rows(min_row=2, values_only=True):

        # Skip completely empty rows
        if not any(value is not None for value in row):
            continue

        book = {
            "book_id": row[0] or "",
            "title": row[1] or "",
            "author": row[2] or "",
            "isbn": row[3] or "",
            "category": row[4] or "",
            "publisher": row[5] or "",
            "year": row[6] or "",
            "copies": row[7] or 0,
            "available": row[8] or 0,
            "shelf": row[9] or "",
            "cover": row[10] or "",
            "date_added": row[11] or ""
        }

        books.append(book)

    workbook.close()

    return books


def add_book(excel_file, book):
    """
    Add a new book to the Excel file.

    Args:
        excel_file: Path to the Excel file.
        book: Dictionary containing book information.
    """

    create_excel_file(excel_file)

    workbook = load_workbook(excel_file)
    sheet = workbook["Books"]

    sheet.append([
        book.get("book_id", ""),
        book.get("title", ""),
        book.get("author", ""),
        book.get("isbn", ""),
        book.get("category", ""),
        book.get("publisher", ""),
        book.get("year", ""),
        book.get("copies", 0),
        book.get("available", 0),
        book.get("shelf", ""),
        book.get("cover", ""),
        book.get("date_added", "")
    ])

    workbook.save(excel_file)
    workbook.close()


def find_book(excel_file, book_id):
    """
    Find a book using its Book ID.

    Returns:
        dict or None
    """

    books = get_all_books(excel_file)

    for book in books:
        if book["book_id"] == book_id:
            return book

    return None


def update_book(excel_file, book_id, updated_book):
    """
    Update an existing book.

    Returns:
        bool: True if updated, False if not found.
    """

    create_excel_file(excel_file)

    workbook = load_workbook(excel_file)
    sheet = workbook["Books"]

    for row in range(2, sheet.max_row + 1):

        current_book_id = sheet.cell(row=row, column=1).value

        if current_book_id == book_id:

            sheet.cell(row=row, column=1).value = book_id
            sheet.cell(row=row, column=2).value = updated_book.get("title", "")
            sheet.cell(row=row, column=3).value = updated_book.get("author", "")
            sheet.cell(row=row, column=4).value = updated_book.get("isbn", "")
            sheet.cell(row=row, column=5).value = updated_book.get("category", "")
            sheet.cell(row=row, column=6).value = updated_book.get("publisher", "")
            sheet.cell(row=row, column=7).value = updated_book.get("year", "")
            sheet.cell(row=row, column=8).value = updated_book.get("copies", 0)
            sheet.cell(row=row, column=9).value = updated_book.get("available", 0)
            sheet.cell(row=row, column=10).value = updated_book.get("shelf", "")
            sheet.cell(row=row, column=11).value = updated_book.get("cover", "")
            
            # Keep original Date Added
            if updated_book.get("date_added"):
                sheet.cell(row=row, column=12).value = updated_book["date_added"]

            workbook.save(excel_file)
            workbook.close()

            return True

    workbook.close()

    return False


def delete_book(excel_file, book_id):
    """
    Delete a book from the Excel file.

    Returns:
        bool: True if deleted, False if not found.
    """

    create_excel_file(excel_file)

    workbook = load_workbook(excel_file)
    sheet = workbook["Books"]

    for row in range(2, sheet.max_row + 1):

        current_book_id = sheet.cell(row=row, column=1).value

        if current_book_id == book_id:

            sheet.delete_rows(row, 1)

            workbook.save(excel_file)
            workbook.close()

            return True

    workbook.close()

    return False


def get_book_count(excel_file):
    """
    Return the total number of books in the Excel file.
    """

    books = get_all_books(excel_file)

    return len(books)


def get_statistics(excel_file):
    """
    Calculate library statistics.

    Returns:
        dict containing:
        - total_books
        - total_copies
        - available_copies
        - borrowed_copies
    """

    books = get_all_books(excel_file)

    total_books = len(books)

    total_copies = sum(
        int(book["copies"])
        for book in books
        if str(book["copies"]).isdigit()
    )

    available_copies = sum(
        int(book["available"])
        for book in books
        if str(book["available"]).isdigit()
    )

    borrowed_copies = total_copies - available_copies

    return {
        "total_books": total_books,
        "total_copies": total_copies,
        "available_copies": available_copies,
        "borrowed_copies": max(borrowed_copies, 0)
    }
