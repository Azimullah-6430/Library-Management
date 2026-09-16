import os
import uuid
import traceback
from datetime import datetime

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    send_from_directory,
    jsonify,
    session,
)

from werkzeug.utils import secure_filename

from openpyxl import Workbook, load_workbook
from PIL import Image


# ============================================================
# CRESCENT COLLEGE LIBRARY MANAGEMENT SYSTEM
# Backend Application
# ============================================================


# ============================================================
# PATH CONFIGURATION
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DATA_FOLDER = os.path.join(
    BASE_DIR,
    "data"
)

UPLOAD_FOLDER = os.path.join(
    BASE_DIR,
    "uploads",
    "book_covers"
)

EXCEL_FILE = os.path.join(
    DATA_FOLDER,
    "library_books.xlsx"
)


# Create required folders
os.makedirs(DATA_FOLDER, exist_ok=True)
os.makedirs(UPLOAD_FOLDER, exist_ok=True)


# ============================================================
# FLASK APPLICATION
# ============================================================

app = Flask(__name__)

app.secret_key = os.environ.get(
    "SECRET_KEY",
    "crescent-library-development-secret-key"
)

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

# Maximum request size: 5 MB
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024


# ============================================================
# EXCEL CONFIGURATION
# ============================================================

HEADERS = [
    "Book ID",
    "Book Title",
    "Author",
    "ISBN",
    "Category",
    "Publisher",
    "Publication Year",
    "Total Copies",
    "Available Copies",
    "Shelf Number",
    "Cover Image",
    "Date Added",
]


# ============================================================
# IMAGE CONFIGURATION
# ============================================================

ALLOWED_IMAGE_EXTENSIONS = {
    "png",
    "jpg",
    "jpeg",
    "webp",
}

MAX_IMAGE_SIZE = 5 * 1024 * 1024


# ============================================================
# EXCEL DATABASE
# ============================================================

def create_empty_excel():
    """
    Create a fresh Excel database with the correct structure.
    """

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Books"

    worksheet.append(HEADERS)

    # Basic formatting
    for cell in worksheet[1]:
        cell.font = cell.font.copy(bold=True)

    column_widths = {
        "A": 15,
        "B": 35,
        "C": 28,
        "D": 22,
        "E": 22,
        "F": 28,
        "G": 20,
        "H": 16,
        "I": 18,
        "J": 18,
        "K": 35,
        "L": 18,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    worksheet.freeze_panes = "A2"

    workbook.save(EXCEL_FILE)
    workbook.close()


def initialize_excel():
    """
    Make sure the Excel database exists and has
    the required Books worksheet.
    """

    try:

        if not os.path.exists(EXCEL_FILE):
            create_empty_excel()
            print("Created new Excel database.")
            return

        workbook = load_workbook(EXCEL_FILE)

        if "Books" not in workbook.sheetnames:

            worksheet = workbook.create_sheet("Books")
            worksheet.append(HEADERS)

            workbook.save(EXCEL_FILE)

        else:

            worksheet = workbook["Books"]

            # Empty workbook/sheet
            if (
                worksheet.max_row == 1
                and worksheet["A1"].value is None
            ):

                worksheet.delete_rows(
                    1,
                    worksheet.max_row
                )

                worksheet.append(HEADERS)

                workbook.save(EXCEL_FILE)

        workbook.close()

    except Exception:

        print("=" * 70)
        print("EXCEL DATABASE ERROR")
        print("=" * 70)

        traceback.print_exc()

        # Close/release the file before backup attempt
        try:
            workbook.close()
        except Exception:
            pass

        # Backup damaged Excel file
        if os.path.exists(EXCEL_FILE):

            backup_file = (
                EXCEL_FILE
                + ".backup_"
                + datetime.now().strftime(
                    "%Y%m%d_%H%M%S"
                )
            )

            try:

                os.rename(
                    EXCEL_FILE,
                    backup_file
                )

                print(
                    f"Damaged Excel file backed up to:\n"
                    f"{backup_file}"
                )

            except Exception:

                print(
                    "Unable to create Excel backup."
                )

        # Create clean database
        create_empty_excel()

        print("Created a fresh Excel database.")
        print("=" * 70)


initialize_excel()


# ============================================================
# EXCEL HELPERS
# ============================================================

def get_workbook():
    """
    Open and return the Excel workbook.
    """

    initialize_excel()

    return load_workbook(
        EXCEL_FILE
    )


def safe_int(value, default=0):
    """
    Safely convert a value into an integer.
    """

    try:

        return int(value)

    except (ValueError, TypeError):

        try:

            return int(float(value))

        except (ValueError, TypeError):

            return default


def normalize_book(row):
    """
    Convert an Excel row into the dictionary structure
    used throughout the application.
    """

    values = list(row)

    while len(values) < 12:
        values.append("")

    total_copies = safe_int(
        values[7],
        0
    )

    available_copies = safe_int(
        values[8],
        0
    )

    # Prevent invalid negative values
    if total_copies < 0:
        total_copies = 0

    if available_copies < 0:
        available_copies = 0

    if available_copies > total_copies:
        available_copies = total_copies

    borrowed_copies = (
        total_copies
        - available_copies
    )

    return {
        "book_id": values[0] or "",
        "title": values[1] or "",
        "author": values[2] or "",
        "isbn": values[3] or "",
        "category": values[4] or "",
        "publisher": values[5] or "",
        "year": values[6] or "",
        "copies": total_copies,
        "available": available_copies,

        # Compatibility aliases
        "available_copies": available_copies,
        "total_copies": total_copies,
        "borrowed": borrowed_copies,
        "borrowed_copies": borrowed_copies,

        "shelf": values[9] or "",
        "cover": values[10] or "",
        "date_added": values[11] or "",
    }


def get_all_books():
    """
    Read all books from the Excel database.
    """

    workbook = None

    try:

        workbook = get_workbook()

        worksheet = workbook["Books"]

        books = []

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            if not any(
                value is not None
                for value in row
            ):
                continue

            books.append(
                normalize_book(row)
            )

        return books

    except Exception:

        traceback.print_exc()

        raise

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass


def find_book(book_id):
    """
    Find a book using its Book ID.
    """

    target_id = str(
        book_id
    ).strip()

    workbook = None

    try:

        workbook = get_workbook()

        worksheet = workbook["Books"]

        for row_number in range(
            2,
            worksheet.max_row + 1
        ):

            current_id = worksheet.cell(
                row=row_number,
                column=1
            ).value

            if (
                str(current_id).strip()
                == target_id
            ):

                row = [
                    worksheet.cell(
                        row=row_number,
                        column=column
                    ).value
                    for column in range(
                        1,
                        13
                    )
                ]

                return normalize_book(row)

        return None

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass


def add_book_to_excel(book):
    """
    Add a new book to Excel.
    """

    workbook = None

    try:

        workbook = get_workbook()

        worksheet = workbook["Books"]

        worksheet.append([
            book["book_id"],
            book["title"],
            book["author"],
            book["isbn"],
            book["category"],
            book["publisher"],
            book["year"],
            book["copies"],
            book["available"],
            book["shelf"],
            book["cover"],
            book["date_added"],
        ])

        workbook.save(EXCEL_FILE)

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass


def update_book_in_excel(book_id, book):
    """
    Update an existing book.
    """

    workbook = None

    try:

        workbook = get_workbook()

        worksheet = workbook["Books"]

        target_id = str(
            book_id
        ).strip()

        for row_number in range(
            2,
            worksheet.max_row + 1
        ):

            current_id = worksheet.cell(
                row=row_number,
                column=1
            ).value

            if (
                str(current_id).strip()
                == target_id
            ):

                values = [
                    book["book_id"],
                    book["title"],
                    book["author"],
                    book["isbn"],
                    book["category"],
                    book["publisher"],
                    book["year"],
                    book["copies"],
                    book["available"],
                    book["shelf"],
                    book["cover"],
                    book["date_added"],
                ]

                for column, value in enumerate(
                    values,
                    start=1
                ):

                    worksheet.cell(
                        row=row_number,
                        column=column
                    ).value = value

                workbook.save(EXCEL_FILE)

                return True

        return False

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass


def delete_book_from_excel(book_id):
    """
    Delete a book from Excel.
    """

    workbook = None

    try:

        workbook = get_workbook()

        worksheet = workbook["Books"]

        target_id = str(
            book_id
        ).strip()

        for row_number in range(
            2,
            worksheet.max_row + 1
        ):

            current_id = worksheet.cell(
                row=row_number,
                column=1
            ).value

            if (
                str(current_id).strip()
                == target_id
            ):

                worksheet.delete_rows(
                    row_number,
                    1
                )

                workbook.save(EXCEL_FILE)

                return True

        return False

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass


# ============================================================
# IMAGE HELPERS
# ============================================================

def is_allowed_image(filename):
    """
    Check whether the image extension is allowed.
    """

    if not filename:
        return False

    if "." not in filename:
        return False

    extension = (
        filename
        .rsplit(".", 1)[1]
        .lower()
    )

    return extension in ALLOWED_IMAGE_EXTENSIONS


def save_cover_image(file):
    """
    Validate and save a book-cover image.

    Returns:
        filename
    """

    if not file:
        return ""


    if not file.filename:
        return ""


    if not is_allowed_image(
        file.filename
    ):

        raise ValueError(
            "Only PNG, JPG, JPEG and WEBP images are allowed."
        )


    # --------------------------------------------------------
    # File size
    # --------------------------------------------------------

    file.seek(
        0,
        os.SEEK_END
    )

    file_size = file.tell()

    file.seek(0)

    if file_size == 0:

        raise ValueError(
            "The uploaded image is empty."
        )

    if file_size > MAX_IMAGE_SIZE:

        raise ValueError(
            "Image size must not exceed 5 MB."
        )


    # --------------------------------------------------------
    # Validate actual image
    # --------------------------------------------------------

    try:

        with Image.open(file) as image:
            image.verify()

        file.seek(0)

    except Exception:

        file.seek(0)

        raise ValueError(
            "The uploaded file is not a valid image."
        )


    # --------------------------------------------------------
    # Generate safe unique filename
    # --------------------------------------------------------

    original_name = secure_filename(
        file.filename
    )

    extension = (
        original_name
        .rsplit(".", 1)[1]
        .lower()
    )

    filename = (
        f"{uuid.uuid4().hex}."
        f"{extension}"
    )

    destination = os.path.join(
        UPLOAD_FOLDER,
        filename
    )


    # --------------------------------------------------------
    # Save image
    # --------------------------------------------------------

    try:

        file.save(destination)

        # Verify saved file
        with Image.open(destination) as image:
            image.verify()

    except Exception:

        if os.path.exists(destination):

            try:
                os.remove(destination)
            except OSError:
                pass

        raise ValueError(
            "Unable to save the book-cover image."
        )


    return filename


def delete_cover_image(filename):
    """
    Delete a stored cover image safely.
    """

    if not filename:
        return

    safe_name = os.path.basename(
        filename
    )

    if safe_name != filename:
        return

    file_path = os.path.join(
        UPLOAD_FOLDER,
        safe_name
    )

    if os.path.isfile(file_path):

        try:
            os.remove(file_path)

        except OSError:

            traceback.print_exc()


# ============================================================
# BOOK ID GENERATOR
# ============================================================

def generate_book_id():
    """
    Generate the next LIBxxxx Book ID.
    """

    books = get_all_books()

    highest_number = 0

    for book in books:

        book_id = str(
            book.get(
                "book_id",
                ""
            )
        ).strip().upper()

        if not book_id.startswith("LIB"):
            continue

        try:

            number = int(
                book_id[3:]
            )

            highest_number = max(
                highest_number,
                number
            )

        except ValueError:

            continue

    return f"LIB{highest_number + 1:04d}"


# ============================================================
# STATISTICS
# ============================================================

def calculate_statistics(books):
    """
    Calculate library statistics.
    """

    total_books = len(books)

    total_copies = sum(
        safe_int(
            book.get(
                "copies",
                0
            )
        )
        for book in books
    )

    available_copies = sum(
        safe_int(
            book.get(
                "available",
                0
            )
        )
        for book in books
    )

    borrowed_copies = (
        total_copies
        - available_copies
    )

    if borrowed_copies < 0:
        borrowed_copies = 0


    return {
        "total_books": total_books,
        "total_copies": total_copies,
        "available_copies": available_copies,
        "borrowed_copies": borrowed_copies,
    }


# ============================================================
# FORM HELPERS
# ============================================================

def get_form_data():
    """
    Collect and clean book form data.
    """

    return {
        "title": request.form.get(
            "title",
            ""
        ).strip(),

        "author": request.form.get(
            "author",
            ""
        ).strip(),

        "isbn": request.form.get(
            "isbn",
            ""
        ).strip(),

        "category": request.form.get(
            "category",
            ""
        ).strip(),

        "publisher": request.form.get(
            "publisher",
            ""
        ).strip(),

        "year": request.form.get(
            "year",
            ""
        ).strip(),

        "copies": request.form.get(
            "copies",
            ""
        ).strip(),

        "shelf": request.form.get(
            "shelf",
            ""
        ).strip(),
    }


def validate_book_data(data):
    """
    Validate book information.
    """

    if not data["title"]:

        return False, (
            "Book title is required."
        )


    if not data["author"]:

        return False, (
            "Author name is required."
        )


    if not data["copies"]:

        return False, (
            "Number of copies is required."
        )


    copies = safe_int(
        data["copies"],
        -1
    )

    if copies < 1:

        return False, (
            "Number of copies must be at least 1."
        )


    if data["year"]:

        year = safe_int(
            data["year"],
            -1
        )

        current_year = datetime.now().year

        if year < 0:

            return False, (
                "Publication year must be a valid number."
            )

        if year > current_year:

            return False, (
                "Publication year cannot be in the future."
            )


    return True, ""


def get_categories(books=None):
    """
    Return unique categories.
    """

    if books is None:
        books = get_all_books()

    categories = set()

    for book in books:

        category = str(
            book.get(
                "category",
                ""
            )
        ).strip()

        if category:
            categories.add(category)

    # Useful default categories
    default_categories = {
        "Computer Science",
        "Information Technology",
        "Engineering",
        "Mathematics",
        "Science",
        "Management",
        "Literature",
        "History",
        "Biography",
        "Other",
    }

    categories.update(
        default_categories
    )

    return sorted(
        categories,
        key=str.lower
    )


# ============================================================
# DASHBOARD
# ============================================================

@app.route("/", endpoint="index")
def dashboard():
    """
    Main library dashboard.

    Endpoint name is intentionally 'index' because
    base.html uses url_for('index').
    """

    try:

        books = get_all_books()

        statistics = calculate_statistics(
            books
        )

        return render_template(
            "dashboard.html",

            books=books,

            statistics=statistics,

            # Template compatibility
            total_books=statistics[
                "total_books"
            ],

            total_copies=statistics[
                "total_copies"
            ],

            available_copies=statistics[
                "available_copies"
            ],

            borrowed_copies=statistics[
                "borrowed_copies"
            ],

            current_date=datetime.now().strftime(
                "%d %B %Y"
            ),
        )

    except Exception:

        print("=" * 70)
        print("DASHBOARD ERROR")
        print("=" * 70)

        traceback.print_exc()

        print("=" * 70)

        return (
            """
            <!DOCTYPE html>
            <html>
            <head>
                <title>Crescent Library - Error</title>
                <style>
                    body {
                        font-family: Arial, sans-serif;
                        background: #f5f7fa;
                        padding: 60px;
                    }

                    .box {
                        max-width: 700px;
                        margin: auto;
                        background: white;
                        padding: 40px;
                        border-radius: 16px;
                        box-shadow:
                            0 10px 30px
                            rgba(0,0,0,0.08);
                    }

                    h1 {
                        color: #c62828;
                    }

                    a {
                        display: inline-block;
                        margin-top: 20px;
                        padding: 12px 22px;
                        background: #111827;
                        color: white;
                        text-decoration: none;
                        border-radius: 8px;
                    }
                </style>
            </head>

            <body>

                <div class="box">

                    <h1>
                        Dashboard Error
                    </h1>

                    <p>
                        The dashboard could not be loaded.
                    </p>

                    <p>
                        Check the terminal for the
                        Python traceback.
                    </p>

                    <a href="/health">
                        Check Application Health
                    </a>

                </div>

            </body>
            </html>
            """,
            500
        )


# ============================================================
# DASHBOARD COMPATIBILITY ROUTE
# ============================================================

@app.route("/dashboard")
def dashboard_alias():
    """
    Compatibility URL for /dashboard.
    """

    return redirect(
        url_for("index")
    )


# ============================================================
# BOOK LIST
# ============================================================

@app.route("/books")
def books():
    """
    Display all books.
    """

    try:

        search_query = request.args.get(
            "search",
            ""
        ).strip()

        all_books = get_all_books()

        if search_query:

            query = search_query.lower()

            all_books = [
                book
                for book in all_books

                if any(
                    query in str(
                        book.get(
                            field,
                            ""
                        )
                    ).lower()

                    for field in [
                        "book_id",
                        "title",
                        "author",
                        "isbn",
                        "category",
                        "publisher",
                        "shelf",
                    ]
                )
            ]


        return render_template(
            "books.html",
            books=all_books,
            search_query=search_query,
        )

    except Exception:

        traceback.print_exc()

        flash(
            "Unable to load books.",
            "danger"
        )

        return redirect(
            url_for("index")
        )


# ============================================================
# ADD BOOK
# ============================================================

@app.route(
    "/add-book",
    methods=["GET", "POST"]
)
def add_book():
    """
    Add a new book.
    """

    books = get_all_books()

    if request.method == "GET":

        return render_template(
            "add_book.html",
            categories=get_categories(books),
        )


    cover_filename = ""

    try:

        data = get_form_data()

        valid, message = validate_book_data(
            data
        )

        if not valid:

            flash(
                message,
                "danger"
            )

            return render_template(
                "add_book.html",
                categories=get_categories(books),
            )


        book_id = generate_book_id()


        # Upload cover
        cover_file = request.files.get(
            "cover"
        )

        if (
            cover_file
            and cover_file.filename
        ):

            cover_filename = save_cover_image(
                cover_file
            )


        copies = safe_int(
            data["copies"]
        )


        book = {

            "book_id": book_id,

            "title": data["title"],

            "author": data["author"],

            "isbn": data["isbn"],

            "category": data["category"],

            "publisher": data["publisher"],

            "year": data["year"],

            "copies": copies,

            "available": copies,

            "shelf": data["shelf"],

            "cover": cover_filename,

            "date_added": datetime.now().strftime(
                "%Y-%m-%d"
            ),
        }


        try:

            add_book_to_excel(book)

        except Exception:

            # Remove uploaded cover if Excel operation fails
            if cover_filename:
                delete_cover_image(
                    cover_filename
                )

            raise


        flash(
            f"Book {book_id} added successfully.",
            "success"
        )

        return redirect(
            url_for("books")
        )


    except ValueError as error:

        flash(
            str(error),
            "danger"
        )

        return render_template(
            "add_book.html",
            categories=get_categories(
                get_all_books()
            ),
        )


    except Exception as error:

        print("ADD BOOK ERROR")

        traceback.print_exc()

        flash(
            f"Unable to add book: {error}",
            "danger"
        )

        return render_template(
            "add_book.html",
            categories=get_categories(
                get_all_books()
            ),
        )


# ============================================================
# BOOK DETAILS
# ============================================================

@app.route(
    "/book/<book_id>"
)
def book_details(book_id):
    """
    Display individual book details.
    """

    try:

        book = find_book(
            book_id
        )

        if not book:

            flash(
                "Book not found.",
                "danger"
            )

            return redirect(
                url_for("books")
            )


        return render_template(
            "book_details.html",
            book=book,
        )


    except Exception:

        print("BOOK DETAILS ERROR")

        traceback.print_exc()

        flash(
            "Unable to load book details.",
            "danger"
        )

        return redirect(
            url_for("books")
        )


# ============================================================
# EDIT BOOK
# ============================================================

@app.route(
    "/edit-book/<book_id>",
    methods=["GET", "POST"]
)
def edit_book(book_id):
    """
    Edit an existing book.
    """

    book = find_book(
        book_id
    )

    if not book:

        flash(
            "Book not found.",
            "danger"
        )

        return redirect(
            url_for("books")
        )


    if request.method == "GET":

        return render_template(
            "edit_book.html",
            book=book,
            categories=get_categories(
                get_all_books()
            ),
        )


    new_cover_filename = ""
    old_cover_filename = book.get(
        "cover",
        ""
    )

    try:

        data = get_form_data()

        valid, message = validate_book_data(
            data
        )

        if not valid:

            flash(
                message,
                "danger"
            )

            return render_template(
                "edit_book.html",
                book=book,
                categories=get_categories(
                    get_all_books()
                ),
            )


        new_total = safe_int(
            data["copies"]
        )

        old_total = safe_int(
            book.get(
                "copies",
                0
            )
        )

        old_available = safe_int(
            book.get(
                "available",
                0
            )
        )


        # Number of currently borrowed copies
        borrowed_copies = (
            old_total
            - old_available
        )

        if borrowed_copies < 0:
            borrowed_copies = 0


        # Cannot reduce total below borrowed count
        if new_total < borrowed_copies:

            flash(
                "Total copies cannot be less than "
                "the number of currently borrowed copies.",
                "danger"
            )

            return render_template(
                "edit_book.html",
                book=book,
                categories=get_categories(
                    get_all_books()
                ),
            )


        new_available = (
            new_total
            - borrowed_copies
        )


        # Existing cover remains unless a new one is uploaded
        cover_filename = old_cover_filename


        new_cover = request.files.get(
            "cover"
        )


        if (
            new_cover
            and new_cover.filename
        ):

            new_cover_filename = save_cover_image(
                new_cover
            )

            cover_filename = new_cover_filename


        updated_book = {

            "book_id": book_id,

            "title": data["title"],

            "author": data["author"],

            "isbn": data["isbn"],

            "category": data["category"],

            "publisher": data["publisher"],

            "year": data["year"],

            "copies": new_total,

            "available": new_available,

            "shelf": data["shelf"],

            "cover": cover_filename,

            "date_added": book.get(
                "date_added",
                datetime.now().strftime(
                    "%Y-%m-%d"
                )
            ),
        }


        # ----------------------------------------------------
        # IMPORTANT:
        # Update Excel FIRST.
        # Delete old cover ONLY after successful update.
        # ----------------------------------------------------

        updated = update_book_in_excel(
            book_id,
            updated_book
        )


        if not updated:

            # Excel failed, remove newly uploaded cover
            if new_cover_filename:
                delete_cover_image(
                    new_cover_filename
                )

            flash(
                "Book could not be updated.",
                "danger"
            )

            return render_template(
                "edit_book.html",
                book=book,
                categories=get_categories(
                    get_all_books()
                ),
            )


        # Excel update successful.
        # Now remove old cover.
        if (
            new_cover_filename
            and old_cover_filename
            and old_cover_filename
            != new_cover_filename
        ):

            delete_cover_image(
                old_cover_filename
            )


        flash(
            "Book updated successfully.",
            "success"
        )


        return redirect(
            url_for(
                "book_details",
                book_id=book_id
            )
        )


    except ValueError as error:

        if new_cover_filename:
            delete_cover_image(
                new_cover_filename
            )

        flash(
            str(error),
            "danger"
        )

        return render_template(
            "edit_book.html",
            book=book,
            categories=get_categories(
                get_all_books()
            ),
        )


    except Exception as error:

        if new_cover_filename:
            delete_cover_image(
                new_cover_filename
            )

        print("EDIT BOOK ERROR")

        traceback.print_exc()

        flash(
            f"Unable to update book: {error}",
            "danger"
        )

        return render_template(
            "edit_book.html",
            book=book,
            categories=get_categories(
                get_all_books()
            ),
        )


# ============================================================
# DELETE BOOK
# ============================================================

@app.route(
    "/delete-book/<book_id>",
    methods=["POST", "GET"]
)
def delete_book(book_id):
    """
    Delete a book.

    Deletion is blocked when copies are borrowed.
    """

    try:

        book = find_book(
            book_id
        )

        if not book:

            flash(
                "Book not found.",
                "danger"
            )

            return redirect(
                url_for("books")
            )


        total = safe_int(
            book.get(
                "copies",
                0
            )
        )

        available = safe_int(
            book.get(
                "available",
                0
            )
        )

        borrowed = (
            total
            - available
        )


        if borrowed > 0:

            flash(
                "This book cannot be deleted because "
                f"{borrowed} copy/copies are currently borrowed.",
                "danger"
            )

            return redirect(
                url_for("books")
            )


        deleted = delete_book_from_excel(
            book_id
        )


        if not deleted:

            flash(
                "Book could not be deleted.",
                "danger"
            )

            return redirect(
                url_for("books")
            )


        # Delete cover only after successful Excel deletion
        if book.get("cover"):

            delete_cover_image(
                book["cover"]
            )


        flash(
            "Book deleted successfully.",
            "success"
        )


    except Exception as error:

        print("DELETE BOOK ERROR")

        traceback.print_exc()

        flash(
            f"Unable to delete book: {error}",
            "danger"
        )


    return redirect(
        url_for("books")
    )


# ============================================================
# CATEGORY BOOKS
# ============================================================

@app.route(
    "/category/<path:category>"
)
def category_books(category):
    """
    Display books belonging to a category.
    """

    try:

        books_list = get_all_books()

        target_category = (
            category
            .strip()
            .lower()
        )

        filtered = [

            book

            for book in books_list

            if (
                str(
                    book.get(
                        "category",
                        ""
                    )
                )
                .strip()
                .lower()
                == target_category
            )
        ]


        return render_template(
            "books.html",
            books=filtered,
            search_query="",
        )


    except Exception:

        traceback.print_exc()

        flash(
            "Unable to load category.",
            "danger"
        )

        return redirect(
            url_for("books")
        )


# ============================================================
# AVAILABLE BOOKS
# ============================================================

@app.route(
    "/available-books"
)
def available_books():
    """
    Display books with available copies.
    """

    books_list = get_all_books()

    available = [

        book

        for book in books_list

        if safe_int(
            book.get(
                "available",
                0
            )
        ) > 0
    ]


    return render_template(
        "books.html",
        books=available,
        search_query="",
    )


# ============================================================
# BORROWED BOOKS
# ============================================================

@app.route(
    "/borrowed-books"
)
def borrowed_books():
    """
    Display books with borrowed copies.
    """

    books_list = get_all_books()

    borrowed = [

        book

        for book in books_list

        if (
            safe_int(
                book.get(
                    "copies",
                    0
                )
            )
            >
            safe_int(
                book.get(
                    "available",
                    0
                )
            )
        )
    ]


    return render_template(
        "books.html",
        books=borrowed,
        search_query="",
    )


# ============================================================
# SEARCH API
# ============================================================

@app.route(
    "/search"
)
def search():
    """
    Search books through JSON API.
    """

    query = request.args.get(
        "q",
        ""
    ).strip().lower()

    books_list = get_all_books()

    if not query:

        return jsonify(
            books_list
        )


    results = [

        book

        for book in books_list

        if any(

            query in str(
                book.get(
                    field,
                    ""
                )
            ).lower()

            for field in [
                "book_id",
                "title",
                "author",
                "isbn",
                "category",
                "publisher",
                "shelf",
            ]
        )
    ]


    return jsonify(
        results
    )


# ============================================================
# STATISTICS API
# ============================================================

@app.route(
    "/api/statistics"
)
def statistics_api():
    """
    Return library statistics as JSON.
    """

    try:

        books_list = get_all_books()

        return jsonify(
            calculate_statistics(
                books_list
            )
        )

    except Exception as error:

        traceback.print_exc()

        return jsonify({
            "error": str(error)
        }), 500


# ============================================================
# BOOKS API
# ============================================================

@app.route(
    "/api/books"
)
def books_api():
    """
    Return all books as JSON.
    """

    try:

        return jsonify(
            get_all_books()
        )

    except Exception as error:

        traceback.print_exc()

        return jsonify({
            "error": str(error)
        }), 500


# ============================================================
# SINGLE BOOK API
# ============================================================

@app.route(
    "/api/books/<book_id>"
)
def book_api(book_id):
    """
    Return one book as JSON.
    """

    book = find_book(
        book_id
    )

    if not book:

        return jsonify({
            "error": "Book not found"
        }), 404


    return jsonify(
        book
    )


# ============================================================
# BOOK COVER ROUTE
# ============================================================

@app.route(
    "/uploads/book_covers/<path:filename>",
    endpoint="uploaded_cover"
)
def uploaded_book_cover(filename):
    """
    Serve uploaded book-cover images.

    Endpoint name is 'uploaded_cover' because
    the HTML templates use url_for('uploaded_cover').
    """

    return send_from_directory(
        UPLOAD_FOLDER,
        filename
    )


# ============================================================
# LOGIN
# ============================================================

@app.route(
    "/login",
    methods=["GET", "POST"]
)
def login():
    """
    Librarian login.

    Demo credentials:
        Username: admin
        Password: admin123
    """

    if request.method == "GET":

        if session.get(
            "logged_in"
        ):

            return redirect(
                url_for("index")
            )

        return render_template(
            "login.html"
        )


    username = request.form.get(
        "username",
        ""
    ).strip()

    password = request.form.get(
        "password",
        ""
    )


    if (
        username == "admin"
        and password == "admin123"
    ):

        session.clear()

        session["logged_in"] = True
        session["username"] = username

        return redirect(
            url_for("index")
        )


    flash(
        "Invalid username or password.",
        "danger"
    )

    return render_template(
        "login.html"
    )


# ============================================================
# LOGOUT
# ============================================================

@app.route(
    "/logout"
)
def logout():

    session.clear()

    return redirect(
        url_for("login")
    )


# ============================================================
# HEALTH CHECK
# ============================================================

@app.route(
    "/health"
)
def health():
    """
    Check whether the application and Excel database
    are working.
    """

    try:

        initialize_excel()

        books = get_all_books()

        return jsonify({

            "status": "healthy",

            "application":
                "Crescent College Library Management System",

            "database":
                "Excel",

            "books":
                len(books),

            "timestamp":
                datetime.now().isoformat(),
        })


    except Exception as error:

        traceback.print_exc()

        return jsonify({

            "status": "error",

            "error":
                str(error),
        }), 500


# ============================================================
# 404 ERROR
# ============================================================

@app.errorhandler(404)
def page_not_found(error):

    return """
    <!DOCTYPE html>

    <html>

    <head>

        <title>
            Crescent Library - 404
        </title>

        <style>

            body {
                font-family: Arial, sans-serif;
                background: #f5f7fa;
                padding: 60px;
            }

            .box {
                max-width: 650px;
                margin: auto;
                background: white;
                padding: 40px;
                border-radius: 16px;
                box-shadow:
                    0 10px 30px
                    rgba(0,0,0,0.08);
            }

            h1 {
                font-size: 52px;
                margin-bottom: 10px;
            }

            a {
                display: inline-block;
                margin-top: 20px;
                padding: 12px 22px;
                background: #111827;
                color: white;
                text-decoration: none;
                border-radius: 8px;
            }

        </style>

    </head>

    <body>

        <div class="box">

            <h1>
                404
            </h1>

            <h2>
                Page Not Found
            </h2>

            <p>
                The requested page does not exist.
            </p>

            <a href="/">
                Go to Dashboard
            </a>

        </div>

    </body>

    </html>
    """, 404


# ============================================================
# FILE TOO LARGE
# ============================================================

@app.errorhandler(413)
def file_too_large(error):

    flash(
        "Uploaded file is too large. Maximum size is 5 MB.",
        "danger"
    )

    return redirect(
        request.referrer
        or url_for("add_book")
    )


# ============================================================
# INTERNAL SERVER ERROR
# ============================================================

@app.errorhandler(500)
def internal_server_error(error):

    print()
    print("=" * 70)
    print("INTERNAL SERVER ERROR")
    print("=" * 70)

    print(
        "Error:",
        error
    )

    traceback.print_exc()

    print("=" * 70)

    return """
    <!DOCTYPE html>

    <html>

    <head>

        <title>
            Crescent Library - Error
        </title>

        <style>

            body {
                font-family: Arial, sans-serif;
                background: #f5f7fa;
                padding: 50px;
            }

            .box {
                max-width: 750px;
                margin: auto;
                background: white;
                padding: 40px;
                border-radius: 16px;
                box-shadow:
                    0 10px 30px
                    rgba(0,0,0,0.08);
            }

            h1 {
                color: #c62828;
            }

            a {
                display: inline-block;
                margin-top: 20px;
                padding: 12px 24px;
                background: #111827;
                color: white;
                text-decoration: none;
                border-radius: 8px;
            }

        </style>

    </head>

    <body>

        <div class="box">

            <h1>
                Internal Server Error
            </h1>

            <p>
                The application encountered an unexpected error.
            </p>

            <p>
                The complete Python traceback is printed
                in the Codespaces terminal.
            </p>

            <a href="/">
                Return to Dashboard
            </a>

        </div>

    </body>

    </html>
    """, 500


# ============================================================
# APPLICATION START
# ============================================================

if __name__ == "__main__":

    print()
    print("=" * 70)
    print("CRESCENT COLLEGE LIBRARY MANAGEMENT SYSTEM")
    print("=" * 70)

    print()
    print("Application : Flask")
    print("Storage     : Excel")
    print()
    print("Excel File:")
    print(EXCEL_FILE)

    print()
    print("Book Covers:")
    print(UPLOAD_FOLDER)

    print()
    print("Dashboard:")
    print("http://127.0.0.1:5000/")

    print()
    print("Health:")
    print("http://127.0.0.1:5000/health")

    print()
    print("Demo Login:")
    print("Username : admin")
    print("Password : admin123")

    print()
    print("=" * 70)

    app.run(
        host="0.0.0.0",
        port=int(
            os.environ.get(
                "PORT",
                5000
            )
        ),
        debug=True,
    )
